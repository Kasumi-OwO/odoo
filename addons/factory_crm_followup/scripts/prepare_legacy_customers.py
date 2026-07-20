"""Convert the specific legacy XLSX export to a stable UTF-8 CSV."""

import argparse
import csv

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "买家昵称",
    "收货地址电话",
    "买家身份",
    "买家等级",
    "所在地-省份",
    "所在地-城市",
    "首次采购日期",
    "最近采购日期",
    "距上次采购（天）",
    "采购次数",
    "累计采购金额（元）",
    "首单是否广告引导",
    "联系名称",
    "联系方式",
    "重要等级",
    "跟进方式",
    "买家需求",
]


parser = argparse.ArgumentParser()
parser.add_argument("source")
parser.add_argument("destination")
args = parser.parse_args()

workbook = load_workbook(args.source, read_only=True, data_only=True)
if len(workbook.worksheets) != 1:
    raise RuntimeError(f"Expected one worksheet, found {len(workbook.worksheets)}")
worksheet = workbook.active
rows = worksheet.iter_rows(values_only=True)
headers = list(next(rows))
if headers != EXPECTED_HEADERS:
    raise RuntimeError(f"Unexpected XLSX headers: {headers!r}")

written = 0
with open(args.destination, "w", encoding="utf-8-sig", newline="") as destination:
    writer = csv.writer(destination)
    writer.writerow(["source_row", *EXPECTED_HEADERS])
    for source_row, values in enumerate(rows, 2):
        if not any(value not in (None, "") for value in values):
            continue
        normalized = [
            value.strftime("%Y-%m-%d %H:%M:%S") if hasattr(value, "strftime") else value
            for value in values
        ]
        writer.writerow([source_row, *normalized])
        written += 1

print(f"PREPARE_RESULT rows={written} destination={args.destination}")
